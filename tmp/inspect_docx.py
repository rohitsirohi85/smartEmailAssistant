from docx import Document
from pathlib import Path
p=Path('output/docx/SmartMail_Assistant_Project_Report.docx')
d=Document(p)
print('sections',len(d.sections),'paragraphs',len(d.paragraphs),'tables',len(d.tables),'inline_shapes',len(d.inline_shapes))
for i,para in enumerate(d.paragraphs[:70]):
    text=para.text.replace('\n',' | ')
    if text.strip(): print(i,repr(text[:180]))
