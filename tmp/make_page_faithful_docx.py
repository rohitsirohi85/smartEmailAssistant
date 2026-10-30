from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

root = Path('tmp/docx_render/source_pages_144dpi')
images = sorted(root.glob('page-*.png'), key=lambda p: int(p.stem.split('-')[-1]))
if len(images) != 78:
    raise SystemExit(f'Expected 78 rendered PDF pages, found {len(images)}')

out = Path('output/docx/SmartMail_Assistant_Project_Report.docx')
doc = Document()
sec = doc.sections[0]
page_w = Inches(8.2680555556)
page_h = Inches(11.6930555556)
sec.page_width = page_w
sec.page_height = page_h
sec.top_margin = Inches(0.05)
sec.bottom_margin = Inches(0.05)
sec.left_margin = Inches(0.05)
sec.right_margin = Inches(0.05)
sec.header_distance = Inches(0.025)
sec.footer_distance = Inches(0.025)

# The source PDF is preserved page-for-page so its university cover, page
# numbering, diagrams, tables, and app screenshots remain visually identical.
image_width = Inches(8.0)
for i, image in enumerate(images):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.keep_together = True
    pf.page_break_after = i < len(images) - 1
    p.add_run().add_picture(str(image.resolve()), width=image_width)

doc.core_properties.title = 'SmartMail Assistant Project Report'
doc.core_properties.subject = 'Page-faithful Word conversion of the approved project report PDF'
doc.core_properties.author = 'Rishi Soni'
doc.save(out)
print(f'Created {out} with {len(images)} full-page images.')
